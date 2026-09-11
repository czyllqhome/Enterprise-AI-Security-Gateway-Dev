import base64
import io
import json
from types import SimpleNamespace
from unittest.mock import Mock

from app.services.file_review_scanner import FileReviewScanner
from app.services.console_service import ConsoleService
from app.services.guardrails.business_sensitive_scanner import (
    BusinessSensitiveCategory,
    BusinessSensitiveResult,
    OpenAICompatibleBusinessSensitiveClient,
)


def model_envelope(result):
    return io.BytesIO(json.dumps({
        "model": "qwen/qwen3.8-flash",
        "choices": [{"message": {"content": json.dumps(result)}}],
    }).encode())


def test_openrouter_image_request_uses_multimodal_payload_and_headers(monkeypatch):
    captured = {}

    def urlopen(req, timeout):
        captured["request"] = req
        captured["timeout"] = timeout
        return model_envelope({
            "contains_business_sensitive": False, "risk_level": "low", "categories": [],
            "summary": "未检测到商务敏感内容。", "confidence": 0.1,
        })

    monkeypatch.setattr("app.services.guardrails.business_sensitive_scanner.request.urlopen", urlopen)
    client = OpenAICompatibleBusinessSensitiveClient(
        api_key="secret", base_url="https://openrouter.ai/api/v1", timeout_seconds=12,
        model="qwen/qwen3.8-flash", provider_label="OpenRouter",
        default_headers={"HTTP-Referer": "http://127.0.0.1:5173", "X-OpenRouter-Title": "Gateway"},
        reasoning_effort="none",
    )
    client.generate_file("review", filename="image.png", mime_type="image/png", content=b"image-bytes")

    payload = json.loads(captured["request"].data)
    image_part = payload["messages"][1]["content"][1]
    assert image_part["type"] == "image_url"
    assert base64.b64decode(image_part["image_url"]["url"].split(",", 1)[1]) == b"image-bytes"
    assert payload["reasoning"] == {"effort": "none"}
    assert captured["request"].headers["Http-referer"] == "http://127.0.0.1:5173"
    assert captured["timeout"] == 12


def test_openrouter_pdf_request_uses_file_content_part(monkeypatch):
    captured = {}

    def urlopen(req, timeout):
        captured["payload"] = json.loads(req.data)
        return model_envelope({
            "contains_business_sensitive": False, "risk_level": "low", "categories": [],
            "summary": "未检测到商务敏感内容。", "confidence": 0.1,
        })

    monkeypatch.setattr("app.services.guardrails.business_sensitive_scanner.request.urlopen", urlopen)
    client = OpenAICompatibleBusinessSensitiveClient(
        api_key="secret", base_url="https://openrouter.ai/api/v1", timeout_seconds=12,
        model="qwen/qwen3.8-flash", provider_label="OpenRouter",
    )
    client.generate_file("review", filename="file.pdf", mime_type="application/pdf", content=b"%PDF-test")
    file_part = captured["payload"]["messages"][1]["content"][1]
    assert file_part["type"] == "file"
    assert file_part["file"]["filename"] == "file.pdf"
    assert file_part["file"]["file_data"].startswith("data:application/pdf;base64,")


def test_medium_multimodal_result_requires_confirmation():
    scanner = FileReviewScanner.__new__(FileReviewScanner)
    scanner.model = "qwen/qwen3.8-flash"
    scanner.runtime = SimpleNamespace(provider="openrouter")
    scanner.business_scanner = SimpleNamespace(scan_file_or_raise=lambda **_: BusinessSensitiveResult(
        contains_business_sensitive=True, risk_level="medium", summary="包含未公开报价。", confidence=0.9,
        categories=[BusinessSensitiveCategory(name="pricing", matched_text="报价 100 万", reason="内部报价")],
    ))
    result = scanner.review_file(filename="quote.pdf", mime_type="application/pdf", content=b"bytes")
    assert result.review_decision == "needs_confirmation"
    assert result.summary == "包含未公开报价。"
    assert result.hits[0].category == "pricing"


def test_office_embedded_image_fails_closed_fast_path(tmp_path):
    from docx import Document
    from PIL import Image
    from app.services.file_extraction_service import FileExtractionService

    image_path = tmp_path / "logo.png"
    Image.new("RGB", (8, 8), "white").save(image_path)
    document_path = tmp_path / "with-image.docx"
    document = Document()
    document.add_paragraph("Public body")
    document.add_picture(str(image_path))
    document.save(document_path)

    extracted = FileExtractionService(SimpleNamespace(image_to_text=lambda _: "")).extract(document_path)
    assert not extracted.coverage_complete
    assert any("save the file as PDF" in issue for issue in extracted.coverage_issues)
    assert extracted.visual_units == []


def test_runtime_update_does_not_reload_all_guardrail_models(monkeypatch):
    reload_events = []

    def guardrail_factory():
        reload_events.append("factory")
        return object()

    guardrail_factory.cache_clear = lambda: reload_events.append("cache_clear")
    monkeypatch.setattr("app.services.console_service.get_guardrail_service", guardrail_factory)

    service = ConsoleService.__new__(ConsoleService)
    service.guardrail_service = object()
    service.setting_service = SimpleNamespace(set_business_sensitive_config=Mock())
    service.get_scanners = lambda: "updated"

    assert service.update_business_sensitive_config("openrouter", "qwen/qwen3.8-flash") == "updated"
    service.setting_service.set_business_sensitive_config.assert_called_once_with(
        provider="openrouter", model="qwen/qwen3.8-flash"
    )
    assert reload_events == []
