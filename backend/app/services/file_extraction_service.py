from __future__ import annotations

import logging
import io
import json
import tempfile
import zipfile
from xml.etree import ElementTree
from pathlib import Path

from ..schemas.file_review_runtime import ExtractedDocument, ExtractedSegmentPayload, VisualReviewUnit
from .file_ocr_service import FileOCRService

logger = logging.getLogger(__name__)


def _normalize_image_metadata(value):
    """Keep metadata reviewable without exposing binary float representation noise."""
    if isinstance(value, float):
        return round(value, 6)
    if isinstance(value, bytes):
        return f"<binary data: {len(value)} bytes>"
    if isinstance(value, dict):
        return {str(key): _normalize_image_metadata(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_normalize_image_metadata(item) for item in value]
    if value is None or isinstance(value, (str, int, bool)):
        return value
    return str(value)


def _serialize_image_metadata(image) -> str:
    metadata = {
        "info": _normalize_image_metadata(dict(image.info)),
        "exif": _normalize_image_metadata(dict(image.getexif())),
    }
    return json.dumps(metadata, ensure_ascii=False, sort_keys=True)


class UnsupportedFileTypeError(Exception):
    pass


class FileExtractionService:
    SUPPORTED_EXTENSIONS = {".docx", ".xlsx", ".pptx", ".pdf", ".png", ".jpg", ".jpeg", ".bmp", ".webp"}

    def __init__(self, ocr_service: FileOCRService | None = None) -> None:
        self.ocr_service = ocr_service or FileOCRService()

    def extract(self, file_path: str | Path) -> ExtractedDocument:
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise UnsupportedFileTypeError(f"Unsupported file type: {extension or 'unknown'}")
        if extension == ".docx":
            return self._mark_office_media_coverage(path, self._extract_docx(path))
        if extension == ".xlsx":
            return self._mark_office_media_coverage(path, self._extract_xlsx(path))
        if extension == ".pptx":
            return self._mark_office_media_coverage(path, self._extract_pptx(path))
        if extension == ".pdf":
            return self._extract_pdf(path)
        return self._extract_image(path)

    def _extract_docx(self, path: Path) -> ExtractedDocument:
        from docx import Document

        document = Document(path)
        segments: list[ExtractedSegmentPayload] = []
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for index, text in enumerate(paragraphs, start=1):
            segments.append(ExtractedSegmentPayload(location=f"Paragraph {index}", text=text))

        for section_index, section in enumerate(document.sections, start=1):
            for area_name, area in (("Header", section.header), ("Footer", section.footer)):
                text = "\n".join(paragraph.text.strip() for paragraph in area.paragraphs if paragraph.text.strip())
                if text:
                    segments.append(ExtractedSegmentPayload(
                        location=f"Section {section_index} {area_name}", text=text, source_kind="header_footer",
                    ))

        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                cell_texts = [cell.text.strip() for cell in row.cells]
                if not any(cell_texts):
                    continue
                text = " | ".join(cell_texts)
                segments.append(
                    ExtractedSegmentPayload(
                        location=f"Table {table_index} Row {row_index}",
                        text=text,
                        source_kind="table",
                    )
                )

        combined = "\n".join(segment.text for segment in segments)
        return ExtractedDocument(summary=f"{len(segments)} extracted segments", plain_text=combined,
                                 segments=segments, coverage_complete=bool(segments),
                                 coverage_issues=[] if segments else ["No extractable Office text was found."])

    def _extract_xlsx(self, path: Path) -> ExtractedDocument:
        from openpyxl import load_workbook

        workbook = load_workbook(filename=path, read_only=True, data_only=False)
        segments: list[ExtractedSegmentPayload] = []
        for sheet in workbook.worksheets:
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [str(value).strip() if value is not None else "" for value in row]
                if not any(values):
                    continue
                text = " | ".join(values)
                segments.append(
                    ExtractedSegmentPayload(
                        location=f"{sheet.title}!Row {row_index}",
                        text=text,
                        sheet_name=sheet.title,
                        source_kind="table",
                    )
                )

        workbook.close()
        combined = "\n".join(segment.text for segment in segments)
        return ExtractedDocument(summary=f"{len(segments)} extracted rows", plain_text=combined,
                                 segments=segments, coverage_complete=bool(segments),
                                 coverage_issues=[] if segments else ["No extractable Office text was found."])

    def _mark_office_media_coverage(self, path: Path, extracted: ExtractedDocument) -> ExtractedDocument:
        """Fast native-text path; unsupported embedded visuals fail closed instead of invoking a vision model."""
        with zipfile.ZipFile(path) as archive:
            media = [entry.filename for entry in archive.infolist()
                     if not entry.is_dir() and "/media/" in entry.filename.lower()]
        if media:
            extracted.coverage_complete = False
            extracted.coverage_issues.append(
                "Embedded Office images are not reviewed in the fast text path; save the file as PDF and upload it again."
            )
        return extracted

    def _audit_office_package(self, path: Path, extracted: ExtractedDocument) -> ExtractedDocument:
        """Include hidden XML text, relationships, metadata and media in the guardrail inventory."""
        from PIL import Image
        issues = []
        with zipfile.ZipFile(path) as archive:
            for entry in archive.infolist():
                if entry.is_dir():
                    continue
                name = entry.filename
                payload = archive.read(entry)
                if name.endswith((".xml", ".rels")):
                    if b"<!DOCTYPE" in payload.upper() or b"<!ENTITY" in payload.upper():
                        issues.append(f"{name}: XML entity declarations are not supported.")
                        continue
                    try:
                        root = ElementTree.fromstring(payload)
                    except ElementTree.ParseError:
                        issues.append(f"{name}: invalid XML.")
                        continue
                    values = []
                    for node in root.iter():
                        if node.tag.rsplit("}", 1)[-1] in {"wsp", "shape", "grpSp"}:
                            issues.append(f"{name}: composed drawing needs rendered visual review.")
                        if node.text and node.text.strip():
                            values.append(node.text.strip())
                        # Attributes include external links, field settings, alternative text and names.
                        values.extend(str(value) for value in node.attrib.values() if str(value).strip())
                    if values:
                        extracted.segments.append(ExtractedSegmentPayload(
                            location=name, text="\n".join(values), source_kind="package_xml"))
                    if any(part in name for part in ("/charts/", "/diagrams/", "/drawings/", "vmlDrawing")):
                        issues.append(f"{name}: rendered Office visual review is required.")
                else:
                    try:
                        with Image.open(io.BytesIO(payload)) as image:
                            frames = getattr(image, "n_frames", 1)
                            if frames > 200:
                                raise ValueError("Too many image frames")
                            extracted.segments.append(ExtractedSegmentPayload(
                                location=name + " metadata", text=_serialize_image_metadata(image)))
                            for index in range(frames):
                                image.seek(index)
                                buffer = io.BytesIO()
                                image.convert("RGB").save(buffer, format="PNG")
                                extracted.visual_units.append(VisualReviewUnit(f"{name} frame {index + 1}", buffer.getvalue()))
                    except Exception:
                        issues.append(f"{name}: embedded binary content requires additional review.")
        if path.suffix.lower() == ".pptx":
            issues.append("Slide layouts require rendered visual review.")
        extracted.coverage_issues = issues
        extracted.coverage_complete = not issues
        extracted.plain_text = "\n".join(segment.text for segment in extracted.segments)
        return extracted

    def _with_office_render(self, path: Path, extracted: ExtractedDocument) -> ExtractedDocument:
        from .office_rendering_service import OfficeRenderingService
        try:
            with tempfile.TemporaryDirectory(prefix="gateway-office-") as directory:
                output = Path(directory)
                rendered = OfficeRenderingService().render_pdf(path, output)
                import fitz
                with fitz.open(rendered) as document:
                    for page_index, page in enumerate(document, 1):
                        text = (page.get_text("text") or "").strip()
                        if text:
                            extracted.segments.append(ExtractedSegmentPayload(
                                location=f"Rendered page {page_index}", text=text, page_number=page_index,
                                source_kind="office_render"))
                        image = page.get_pixmap(dpi=160).tobytes("png")
                        extracted.visual_units.append(VisualReviewUnit(f"Rendered page {page_index}", image))
            # Rendering supplies the visual surface for package drawings/layouts. Other binary issues remain.
            extracted.coverage_issues = [issue for issue in extracted.coverage_issues if not any(
                marker in issue for marker in ("rendered Office visual review", "composed drawing", "Slide layouts"))]
            extracted.coverage_complete = not extracted.coverage_issues
            extracted.plain_text = "\n".join(segment.text for segment in extracted.segments)
        except Exception as exc:
            extracted.coverage_complete = False
            extracted.coverage_issues.append(f"Office rendering unavailable: {exc}")
        return extracted

    def _extract_pptx(self, path: Path) -> ExtractedDocument:
        from pptx import Presentation

        presentation = Presentation(path)
        segments: list[ExtractedSegmentPayload] = []

        for slide_index, slide in enumerate(presentation.slides, start=1):
            for shape_index, shape in enumerate(slide.shapes, start=1):
                if getattr(shape, "has_text_frame", False):
                    text = "\n".join(paragraph.text.strip() for paragraph in shape.text_frame.paragraphs if paragraph.text.strip())
                    if text:
                        segments.append(
                            ExtractedSegmentPayload(
                                location=f"Slide {slide_index} Shape {shape_index}",
                                text=text,
                                slide_number=slide_index,
                            )
                        )
                if getattr(shape, "has_table", False):
                    for row_index, row in enumerate(shape.table.rows, start=1):
                        text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if text:
                            segments.append(ExtractedSegmentPayload(
                                location=f"Slide {slide_index} Table {shape_index} Row {row_index}", text=text,
                                slide_number=slide_index, source_kind="table",
                            ))

        combined = "\n".join(segment.text for segment in segments)
        return ExtractedDocument(summary=f"{len(segments)} extracted slide items", plain_text=combined,
                                 segments=segments, coverage_complete=bool(segments),
                                 coverage_issues=[] if segments else ["No extractable Office text was found."])

    def _extract_pdf(self, path: Path) -> ExtractedDocument:
        import fitz

        segments: list[ExtractedSegmentPayload] = []
        issues: list[str] = []
        with fitz.open(path) as document:
            if document.needs_pass:
                raise ValueError("Encrypted PDFs cannot be reviewed.")
            if document.embfile_count() or document.get_ocgs():
                issues.append("PDF embedded files or optional layers require additional review.")
            for xref in range(1, document.xref_length()):
                obj = document.xref_object(xref)
                if any(marker in obj for marker in ("/JavaScript", "/JS", "/Launch", "/RichMedia", "/XFA", "/Filespec")):
                    issues.append(f"PDF object {xref} contains unsupported active or embedded content.")
            metadata = "\n".join(f"{key}: {value}" for key, value in document.metadata.items() if value)
            metadata += "\n" + (document.get_xml_metadata() or "")
            if metadata.strip():
                segments.append(ExtractedSegmentPayload(location="PDF metadata", text=metadata))
            for page_index, page in enumerate(document, start=1):
                text = (page.get_text("text") or "").strip()
                if text:
                    segments.append(ExtractedSegmentPayload(location=f"Page {page_index}", text=text, page_number=page_index))
                links = "\n".join(str(link.get("uri") or link.get("file") or "") for link in page.get_links())
                if links.strip():
                    segments.append(ExtractedSegmentPayload(location=f"Page {page_index} links", text=links))
                visual = bool(page.get_images(full=True) or page.get_drawings())
                if list(page.annots() or []) or list(page.widgets() or []):
                    issues.append(f"Page {page_index}: annotations or form fields require additional review.")
                if visual or not text:
                    # Mixed pages must not skip OCR merely because some native text exists.
                    try:
                        pixmap = page.get_pixmap(dpi=160)
                        image_bytes = pixmap.tobytes("png")
                        ocr_text = self._ocr_bytes(image_bytes, suffix=".png")
                        if ocr_text:
                            segments.append(ExtractedSegmentPayload(location=f"Page {page_index} OCR", text=ocr_text,
                                                                    page_number=page_index, source_kind="image_ocr"))
                    except Exception:
                        issues.append(f"Page {page_index}: OCR failed.")
        combined = "\n".join(segment.text for segment in segments)
        return ExtractedDocument(summary=f"{len(segments)} extracted PDF items", plain_text=combined,
                                 segments=segments, coverage_complete=not issues, coverage_issues=issues,
                                 visual_units=[])

    def _extract_image(self, path: Path) -> ExtractedDocument:
        from PIL import Image
        text = self.ocr_service.image_to_text(path)
        segment = ExtractedSegmentPayload(location="Image", text=text, source_kind="image_ocr")
        segments = [segment] if text else []
        with Image.open(path) as image:
            frames = getattr(image, "n_frames", 1)
            if frames > 200:
                raise ValueError("Image exceeds the 200-frame review budget.")
        issues = [] if text and frames == 1 else (["Local OCR found no readable image text."] if not text
                                                   else ["Animated or multi-frame images are not fully covered."])
        return ExtractedDocument(summary="Image text extracted with one local OCR pass", plain_text=text,
                                 segments=segments, coverage_complete=not issues,
                                 coverage_issues=issues, visual_units=[])

    def _ocr_bytes(self, payload: bytes, *, suffix: str) -> str:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as handle:
            handle.write(payload)
            temp_path = Path(handle.name)
        try:
            return self.ocr_service.image_to_text(temp_path)
        finally:
            temp_path.unlink(missing_ok=True)
