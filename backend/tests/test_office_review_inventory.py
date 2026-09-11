from types import SimpleNamespace

from app.services.file_extraction_service import FileExtractionService


def extractor():
    return FileExtractionService(SimpleNamespace(image_to_text=lambda _: ""))


def test_word_headers_and_empty_table_cells_are_in_review_inventory(tmp_path):
    from docx import Document
    path = tmp_path / "file.docx"
    document = Document()
    document.add_paragraph("Public body")
    document.sections[0].header.paragraphs[0].text = "HEADER_PRIVATE_MARKER"
    cells = document.add_table(rows=1, cols=3).rows[0].cells
    cells[0].text = "Supplier"
    cells[2].text = "Amount"
    document.save(path)
    result = extractor().extract(path)
    assert "HEADER_PRIVATE_MARKER" in result.plain_text
    assert any(segment.text == "Supplier |  | Amount" for segment in result.segments)
    assert result.visual_units == []
    assert result.coverage_complete


def test_workbook_hidden_rows_and_formulas_are_reviewed(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "file.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Item"
    sheet["C1"] = "Amount"
    sheet["A2"] = "HIDDEN_PRIVATE_MARKER"
    sheet.row_dimensions[2].hidden = True
    sheet["C2"] = "=SUM(1,2)"
    workbook.save(path)
    workbook.close()
    result = extractor().extract(path)
    assert "HIDDEN_PRIVATE_MARKER" in result.plain_text
    assert "SUM(1,2)" in result.plain_text
    assert any(segment.text == "Item |  | Amount" for segment in result.segments)
    assert result.coverage_complete


def test_slide_native_text_does_not_require_visual_rendering(tmp_path):
    from pptx import Presentation
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Public slide"
    presentation.save(path)
    result = extractor().extract(path)
    assert "Public slide" in result.plain_text
    assert result.coverage_complete
    assert result.visual_units == []
